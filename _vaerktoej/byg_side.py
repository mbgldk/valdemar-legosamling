"""Bygger hjemmesiden (index.html + img/) ud fra Legosamling.xlsx og mappen Produktbilleder.

Brug: python3 _vaerktoej/byg_side.py <Legosamling.xlsx> <Produktbilleder-mappe> <repo-mappe>
"""
import sys, os, re, json, glob, datetime
from openpyxl import load_workbook
from PIL import Image

xlsx, imgdir, repo = sys.argv[1], sys.argv[2], sys.argv[3]
here = os.path.dirname(os.path.abspath(__file__))

# ---- læs arket efter kolonneoverskrifter (række 4) ----
ws = load_workbook(xlsx, data_only=False)["Samling"]
H = {str(ws.cell(4, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
def val(r, name):
    c = H.get(name)
    return ws.cell(r, c).value if c else None

imgs = {os.path.basename(p)[:5]: p for p in glob.glob(os.path.join(imgdir, "*"))
        if re.match(r"\d{5} - ", os.path.basename(p))}

outimg = os.path.join(repo, "img")
os.makedirs(outimg, exist_ok=True)
sets, used = [], set()
for r in range(5, ws.max_row + 1):
    nr = val(r, "Sætnr")
    if not nr:
        continue
    nr = str(nr).strip()
    dele = val(r, "Antal dele")
    aar = val(r, "År")
    x = {"nr": nr, "navn": val(r, "Navn") or nr, "tema": val(r, "Tema") or "Andet",
         "aar": int(aar) if isinstance(aar, (int, float)) else "",
         "dele": int(dele) if isinstance(dele, (int, float)) else 0, "img": "",
         # tom PDF 1 = LEGO har ingen vejledning til sættet (se Note i arket)
         "vejl": bool(str(val(r, "PDF 1") or "").strip())}
    p = imgs.get(nr)
    if p:
        fn = f"{nr}.webp"
        im = Image.open(p)
        im = im.convert("RGBA") if im.mode in ("RGBA", "LA", "P") else im.convert("RGB")
        im.thumbnail((600, 450))
        im.save(os.path.join(outimg, fn), "WEBP", quality=82)
        x["img"] = "img/" + fn
        used.add(fn)
    sets.append(x)

# fjern billeder af sæt, der ikke længere er i arket
for f in os.listdir(outimg):
    if f not in used:
        os.remove(os.path.join(outimg, f))

maaneder = ["januar", "februar", "marts", "april", "maj", "juni", "juli", "august",
            "september", "oktober", "november", "december"]
d = datetime.date.today()
dato = f"{d.day}. {maaneder[d.month - 1]} {d.year}"

html = open(os.path.join(here, "skabelon.html"), encoding="utf-8").read()
html = html.replace("__SETS__", json.dumps(sets, ensure_ascii=False)).replace("__DATO__", dato)
open(os.path.join(repo, "index.html"), "w", encoding="utf-8").write(html)
open(os.path.join(repo, "CNAME"), "w").write("lego.mbiegel.dk\n")
open(os.path.join(repo, "robots.txt"), "w").write("User-agent: *\nDisallow: /\n")
print(f"{len(sets)} sæt, {len(used)} billeder, opdateret {dato}")
